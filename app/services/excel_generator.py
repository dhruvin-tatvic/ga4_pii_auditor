import os
import tempfile
import openpyxl
from openpyxl.styles import Font, PatternFill

class ExcelGenerator:
    def __init__(self, temp_dir=None):
        self.temp_dir = temp_dir or tempfile.gettempdir()
        os.makedirs(self.temp_dir, exist_ok=True)
    
    def remove_event_prefix(self, value):
        """Remove event prefixes like 'customEvent:' from dimension names"""
        if isinstance(value, str) and ':' in value:
            return value.split(':', 1)[1]
        return value

    def generate_excel_report(self, client_name, property_id, start_date, end_date, leaks):
        """
        Generates an Excel report of the PII leaks.
        Returns the absolute path to the generated .xlsx file.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "PII Leaks Report"

        # Write header
        headers = ["Dimension", "Flagged Value", "PII Type"]
        ws.append(headers)
        
        # Style header
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        
        for col_num, cell in enumerate(ws[1], 1):
            cell.fill = header_fill
            cell.font = header_font
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = 30

        # Write data
        for leak in leaks:
            ws.append([
                self.remove_event_prefix(leak.get("dimension", "Unknown")),
                leak.get("flagged_value", "Unknown"),
                leak.get("type", "Unknown")
            ])

        # Generate a safe filename
        safe_client_name = "".join([c if c.isalnum() else "_" for c in client_name])
        filename = f"PII_Leaks_{safe_client_name}_{property_id}.xlsx"
        filepath = os.path.join(self.temp_dir, filename)
        
        wb.save(filepath)
        return filepath
